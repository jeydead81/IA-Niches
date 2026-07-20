from openpyxl import load_workbook

from fiction_validation import export_validation
from models import EnrichedBook, TropeClassification


def _fichier_corrige(tmp_path):
    """xlsx exporté puis corrigé par Baptiste, comme après un vrai passage manuel."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1",
                              tropes=["metier_gourmand"], decor="village_breton")]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    col = [c.value for c in ws[1]].index("tropes_ok") + 1
    ws.cell(row=2, column=col).value = "metier_gourmand"     # Baptiste confirme
    wb.save(str(p))
    return p


def test_cli_imprime_le_taux_et_la_porte(tmp_path, capsys):
    """Le CLI est le seul moyen de produire le rapport sans re-classifier (non
    déterministe) — cf. A1. Preuve minimale : il tourne et affiche taux + porte."""
    from validate_classifier import main

    p = _fichier_corrige(tmp_path)
    main([str(p)])
    sortie = capsys.readouterr().out
    assert "100" in sortie or "1.0" in sortie          # taux d'accord (1 livre, accord total)
    assert "franchie" in sortie.lower() or "80" in sortie
