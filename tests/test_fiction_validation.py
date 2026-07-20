from openpyxl import load_workbook

from fiction_validation import (accord, agreement_report, export_validation, jaccard,
                                load_corrections)
from models import EnrichedBook, TropeClassification


def test_jaccard_deux_vides_est_un_accord():
    assert jaccard(set(), set()) == 1.0


def test_livre_en_accord_exige_les_trois_criteres():
    ia = TropeClassification(asin="A", taxonomy_version="v", tropes=["a", "b"], decor="d")
    hu = TropeClassification(asin="A", taxonomy_version="v", tropes=["a"], decor="d")
    assert accord(ia, hu) is True                        # jaccard 0.5, décor ok
    hu2 = hu.model_copy(update={"decor": "autre"})
    assert accord(ia, hu2) is False                      # décor diverge -> désaccord


def test_rapport_donne_le_taux_et_la_porte():
    paires = [(TropeClassification(asin=str(i), taxonomy_version="v", tropes=["a"]),
               TropeClassification(asin=str(i), taxonomy_version="v",
                                   tropes=["a"] if i < 8 else ["z"])) for i in range(10)]
    r = agreement_report(paires)
    assert r.n_livres == 10 and r.n_accord == 8
    assert r.taux == 0.8 and r.porte_franchie is True


def test_rapport_pointe_les_cles_litigieuses():
    """Sous la porte, il faut savoir QUELLE clé pose problème — sinon on ne peut rien corriger."""
    paires = [(TropeClassification(asin="1", taxonomy_version="v", tropes=["flou"]),
               TropeClassification(asin="1", taxonomy_version="v", tropes=["net"]))]
    r = agreement_report(paires)
    assert r.porte_franchie is False
    assert r.cles_litigieuses["flou"]["ia_seule"] == 1
    assert r.cles_litigieuses["net"]["humain_seul"] == 1


def test_blurb_tronque_a_2000_caracteres_avec_marque_explicite(tmp_path):
    """Mesuré sur les fixtures live : à 600 caractères, l'humain ne voyait que 44 % du
    texte lu par l'IA, sans aucune marque de troncature — le désaccord mesurait alors une
    asymétrie d'information, pas un vrai désaccord de lecture."""
    long_blurb = "x" * 2500
    livres = [EnrichedBook(asin="A1", title="T", blurb=long_blurb)]
    p = tmp_path / "v.xlsx"
    export_validation(livres, [], "cosy_mystery", p)
    wb = load_workbook(str(p))
    cell = wb["Validation"].cell(row=2, column=4).value
    assert cell.startswith("x" * 2000)
    assert not cell.startswith("x" * 2001)
    assert "tronqué" in cell


def test_blurb_court_non_tronque_sans_marque(tmp_path):
    livres = [EnrichedBook(asin="A1", title="T", blurb="un blurb court")]
    p = tmp_path / "v.xlsx"
    export_validation(livres, [], "cosy_mystery", p)
    wb = load_workbook(str(p))
    cell = wb["Validation"].cell(row=2, column=4).value
    assert cell == "un blurb court"


def _idx(ws):
    return {c.value: i for i, c in enumerate(ws[1], start=1)}


def test_livre_non_classe_a_un_statut_explicite_et_cellules_vides(tmp_path):
    """Un livre absent de la réponse du LLM ne doit pas ressembler à « rien trouvé » :
    sans statut explicite, tropes vides + est_roman=True + confiance=0 est indiscernable
    d'une vraie lecture qui n'a identifié aucun trope."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="un blurb")]
    p = tmp_path / "v.xlsx"
    export_validation(livres, [], "cosy_mystery", p)      # aucune classification IA fournie
    ws = load_workbook(str(p))["Validation"]
    idx = _idx(ws)
    row = [c.value for c in ws[2]]
    assert row[idx["statut_ia"] - 1] == "NON CLASSÉ"
    assert row[idx["tropes_ia"] - 1] in (None, "")
    assert row[idx["decor_ia"] - 1] in (None, "")
    assert row[idx["est_roman_ia"] - 1] in (None, "")
    assert row[idx["confiance"] - 1] in (None, "")


def test_livre_classe_a_un_statut_classe(tmp_path):
    livres = [EnrichedBook(asin="A1", title="T", blurb="un blurb")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    ws = load_workbook(str(p))["Validation"]
    row = [c.value for c in ws[2]]
    assert row[_idx(ws)["statut_ia"] - 1] == "classé"


def test_taxonomy_version_ecrite_dans_une_cellule_visible(tmp_path):
    """`wb.properties.keywords` devient vide si Excel efface les propriétés du classeur à
    l'enregistrement (« enregistrer sous ») : la version doit aussi vivre dans une cellule
    visible pour rester lisible dans ce cas."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    ws2 = load_workbook(str(p))["Clés autorisées"]
    assert ws2["B1"].value == "fr_v1"


def test_taxonomy_version_relue_depuis_la_cellule_meme_si_les_proprietes_sont_effacees(tmp_path):
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    wb.properties.keywords = ""                 # simule Excel qui efface les propriétés
    wb.save(str(p))
    # une ligne corrigée doit reconstruire taxonomy_version="fr_v1" via la cellule visible,
    # pas via les propriétés (vides)
    ws = wb["Validation"]
    ws.cell(row=2, column=[c.value for c in ws[1]].index("tropes_ok") + 1).value = "metier_gourmand"
    wb.save(str(p))
    corrections = load_corrections(p)
    assert corrections[0].taxonomy_version == "fr_v1"


def test_taxonomy_version_introuvable_leve(tmp_path):
    """Mesurer contre une taxonomie inconnue n'a pas de sens : si ni la cellule visible ni
    les propriétés du classeur ne portent la version, il faut lever plutôt que de mesurer
    silencieusement contre une chaîne vide."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    wb.properties.keywords = ""
    wb["Clés autorisées"]["B1"].value = None
    ws = wb["Validation"]
    ws.cell(row=2, column=[c.value for c in ws[1]].index("tropes_ok") + 1).value = "metier_gourmand"
    wb.save(str(p))
    import pytest
    with pytest.raises(ValueError):
        load_corrections(p)


def _col(ws, header: str) -> int:
    return [c.value for c in ws[1]].index(header) + 1


def test_cles_humaines_normalisees_avant_comparaison(tmp_path):
    """" Metier_Gourmand " ne doit pas compter comme un désaccord contre metier_gourmand :
    seule la casse/les espaces diffèrent, pas le sens — sans normalisation ça accuse à tort
    la taxonomie d'une clé « litigieuse » qui n'existe même pas."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "tropes_ok")).value = " Metier_Gourmand "
    wb.save(str(p))
    out = load_corrections(p)
    assert out[0].tropes == ["metier_gourmand"]


def test_cle_hors_taxonomie_saisie_est_une_faute_de_saisie_pas_un_trope(tmp_path):
    """Une clé saisie par Baptiste qui n'existe pas dans la taxonomie du sous-genre est une
    FAUTE DE SAISIE (à corriger côté Excel), pas une observation qui accuse la taxonomie."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "tropes_ok")).value = "cle_totalement_inventee"
    wb.save(str(p))
    out = load_corrections(p)
    assert out[0].tropes == []
    assert "cle_totalement_inventee" in out[0].fautes_saisie


def test_colonne_ok_vide_sur_ligne_corrigee_retombe_sur_l_ia(tmp_path):
    """Le mode d'emploi dit « corrige uniquement ce qui te semble faux » : sur une ligne
    corrigée, une colonne *_ok laissée vide doit reprendre l'étiquette IA de la même ligne,
    jamais retomber sur le vide — sinon contester le seul décor efface aussi les tropes et
    fait chuter le taux à 0."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1",
                              tropes=["metier_gourmand"], decor="village_breton")]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "decor_ok")).value = "provence"   # ne conteste QUE le décor
    wb.save(str(p))
    out = load_corrections(p)
    assert out[0].decor == "provence"
    assert out[0].tropes == ["metier_gourmand"]      # pas effacé par la correction du décor


def test_notes_seule_ne_vaut_pas_correction(tmp_path):
    """Une note « pas sûr » ne doit pas fabriquer une étiquette humaine : seules les
    colonnes *_ok comptent comme un avis (c'est déjà ce que dit la docstring de
    load_corrections, mais le code comptait aussi `notes` dans la condition)."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", tropes=["metier_gourmand"])]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "notes")).value = "pas sûr"     # aucune colonne *_ok touchée
    wb.save(str(p))
    assert load_corrections(p) == []


def test_est_roman_ok_reconnait_les_formes_negatives_usuelles(tmp_path):
    """Avant le correctif, seule la forme exacte "non" (parmi false/faux/0/non) était
    reconnue : "n" (paire o/n) tombait dans le else -> True, sur le champ même qui sert à
    écarter les jeux et cahiers de coloriage."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", est_roman=True)]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "est_roman_ok")).value = "n"
    wb.save(str(p))
    out = load_corrections(p)
    assert out[0].est_roman is False


def test_est_roman_ok_valeur_libre_non_reconnue_retombe_sur_l_ia_et_est_signalee(tmp_path):
    """Une valeur libre non reconnue ("jeu", "pas un roman"...) ne doit pas s'inverser en
    True par défaut. Elle doit être traitée comme une colonne NON corrigée (retombe sur
    l'IA de la même ligne) et être signalée plutôt que devinée en silence."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1", est_roman=False)]
    p = tmp_path / "v.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    wb = load_workbook(str(p))
    ws = wb["Validation"]
    ws.cell(row=2, column=_col(ws, "est_roman_ok")).value = "jeu"
    wb.save(str(p))
    out = load_corrections(p)
    assert out[0].est_roman is False        # retombe sur l'IA (False), pas un True deviné
    assert any("jeu" in f for f in out[0].fautes_saisie)


def test_agreement_report_leve_si_les_asin_de_la_paire_divergent():
    """Erreur de programmation, pas une donnée : une paire mal construite doit lever,
    jamais fausser silencieusement le taux mesuré."""
    import pytest
    ia = TropeClassification(asin="A", taxonomy_version="v")
    humain = TropeClassification(asin="B", taxonomy_version="v")
    with pytest.raises(ValueError):
        agreement_report([(ia, humain)])


def test_livre_jamais_classe_exclu_du_taux_et_compte_a_part():
    """Prouve l'exploit d'origine : 3 livres jamais classés par l'IA (seul est_roman_ok
    rempli côté humain) comptaient comme un accord parfait (tropes vides des deux côtés,
    décor None, est_roman par défaut identique) -> taux 1.0, porte franchie, sans qu'aucune
    vraie lecture n'ait été comparée. `ia=None` signale « jamais classé »."""
    paires = [(None, TropeClassification(asin=str(i), taxonomy_version="v", est_roman=True))
              for i in range(3)]
    r = agreement_report(paires)
    assert r.n_non_classes == 3
    assert r.n_livres == 3
    assert r.taux == 0.0                  # rien de mesurable -> pessimiste, pas 100 %
    assert r.porte_franchie is False


def test_export_puis_relecture_conserve_les_etiquettes(tmp_path):
    """Aller-retour Excel : ce que Baptiste corrige doit revenir tel quel."""
    livres = [EnrichedBook(asin="A1", title="Titre", blurb="Un blurb.")]
    ia = [TropeClassification(asin="A1", taxonomy_version="fr_v1",
                              tropes=["metier_gourmand"], decor="village_breton")]
    p = tmp_path / "valid.xlsx"
    export_validation(livres, ia, "cosy_mystery", p)
    assert p.exists()
    relu = load_corrections(p)            # sans correction humaine -> colonnes vides
    assert relu == []                     # rien de corrigé = rien à comparer
