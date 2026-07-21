from build_validation_set import build_set
from models import EnrichedBook, FictionNiche, FictionShelf, TropeClassification


def _niche(query):
    # sous_genre doit être une clé RÉELLE de fiction_taxonomy : export_validation (M4-3)
    # y appelle valid_keys() pour la feuille "Clés autorisées", qui lève KeyError sur une
    # clé inventée. Le plan proposait "cosy mystery a"/"b" comme simples libellés de niche ;
    # adapté ici en deux TRIOS distincts (query différente) du même sous-genre réel — ce que
    # produit d'ailleurs le CLI en usage normal (plusieurs trios pour UN sous-genre).
    return FictionNiche(sous_genre="cosy_mystery", query=query)


def test_assemble_les_rayons_et_deduplique(tmp_path):
    livres = [EnrichedBook(asin="A1", title="T", blurb="b"),
              EnrichedBook(asin="A2", title="T", blurb="b")]

    def faux_shelf(niche, **kw):
        return FictionShelf(niche=niche, search_param="x", books=livres,
                            asins_demandes=2)

    out = tmp_path / "v.xlsx"
    n = build_set([_niche("cosy mystery a"), _niche("cosy mystery b")], out,
                  fetch_shelf=faux_shelf,
                  classify=lambda bks, sg, **kw: [
                      TropeClassification(asin=b.asin, taxonomy_version="fr_v1") for b in bks])
    assert n == 2                 # A1/A2 vus deux fois -> comptés une seule
    assert out.exists()


def test_rayon_en_echec_nempeche_pas_lexport_des_rayons_deja_payes(tmp_path):
    """Un rayon qui lève ne doit pas faire perdre les rayons DÉJÀ payés (ni empêcher tout
    export) : chaque rayon est isolé (try/except)."""
    livres_ok = [EnrichedBook(asin="A1", title="T", blurb="b")]

    def shelf_instable(niche, **kw):
        if niche.query == "casse":
            raise RuntimeError("panne réseau")
        return FictionShelf(niche=niche, search_param="x", books=livres_ok, asins_demandes=1)

    out = tmp_path / "v.xlsx"
    n = build_set([_niche("ok"), _niche("casse")], out, fetch_shelf=shelf_instable,
                  classify=lambda bks, sg, **kw: [])
    assert n == 1                 # le rayon "ok" est quand même exporté
    assert out.exists()


def test_echecs_signales_a_la_fin_du_run(tmp_path, capsys):
    """FictionShelf.n_echecs (échecs d'enrichissement au sein d'un rayon récupéré) n'était
    jamais consulté ; le nombre de rayons entièrement en échec non plus. Un échec silencieux
    se lirait comme un set complet et sain."""
    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]

    def shelf_partiel(niche, **kw):
        if niche.query == "casse":
            raise RuntimeError("panne réseau")
        return FictionShelf(niche=niche, search_param="x", books=livres, asins_demandes=3,
                            n_echecs=2)

    out = tmp_path / "v.xlsx"
    build_set([_niche("ok"), _niche("casse")], out, fetch_shelf=shelf_partiel,
             classify=lambda bks, sg, **kw: [])
    sortie = capsys.readouterr().out
    assert "rayon" in sortie.lower() and "échec" in sortie.lower()
    assert "1" in sortie          # 1 rayon entièrement en échec ("casse")
    assert "2" in sortie          # 2 échecs d'enrichissement dans le rayon "ok"


def test_plafond_50_livres(tmp_path):
    """Le set est dimensionné pour ~50 livres (protocole M4-3) : un plafond évite de payer
    et de classer inutilement au-delà."""
    livres = [EnrichedBook(asin=f"A{i}", title="T", blurb="b") for i in range(80)]

    def faux_shelf(niche, **kw):
        return FictionShelf(niche=niche, search_param="x", books=livres, asins_demandes=80)

    out = tmp_path / "v.xlsx"
    n = build_set([_niche("large")], out, fetch_shelf=faux_shelf,
                  classify=lambda bks, sg, **kw: [])
    assert n == 50


def test_cles_autorisees_couvre_tous_les_sous_genres_du_set(tmp_path):
    """Un set peut mélanger plusieurs sous-genres : la feuille « Clés autorisées » doit
    couvrir chacun d'eux, pas seulement celui de la première niche."""
    livres_cosy = [EnrichedBook(asin="A1", title="T", blurb="b")]
    livres_thriller = [EnrichedBook(asin="B1", title="T", blurb="b")]

    def faux_shelf(niche, **kw):
        livres = livres_cosy if niche.sous_genre == "cosy_mystery" else livres_thriller
        return FictionShelf(niche=niche, search_param="x", books=livres, asins_demandes=1)

    niches = [FictionNiche(sous_genre="cosy_mystery", query="a"),
             FictionNiche(sous_genre="thriller_psychologique", query="b")]
    out = tmp_path / "v.xlsx"
    build_set(niches, out, fetch_shelf=faux_shelf, classify=lambda bks, sg, **kw: [])

    from openpyxl import load_workbook

    from fiction_taxonomy import valid_keys
    valeurs = {c[0].value for c in load_workbook(str(out))["Clés autorisées"].iter_rows()
              if c[0].value}
    assert "metier_gourmand" in valeurs                          # cosy_mystery
    tropes_thriller, _ = valid_keys("thriller_psychologique")
    assert any(t in valeurs for t in tropes_thriller)             # thriller_psychologique


def test_le_cout_llm_de_la_classification_est_compte(tmp_path):
    """build_set imprime « coût réel du run » : sans brancher on_usage sur le CostTracker,
    la moitié LLM de la dépense est absente et le chiffre affiché est faux (§10). Vu en
    conditions réelles : 50 livres classés, llm_usd = 0 au rapport."""
    from cost_tracker import CostTracker

    livres = [EnrichedBook(asin="A1", title="T", blurb="b")]

    def faux_shelf(niche, **kw):
        return FictionShelf(niche=niche, search_param="x", books=livres, asins_demandes=1)

    def faux_classify(bks, sg, on_usage=None, **kw):
        if on_usage:                       # le classifieur réel appelle ce callback
            on_usage(1000, 200, "claude-sonnet-5")
        return [TropeClassification(asin=b.asin, taxonomy_version="fr_v1") for b in bks]

    cost = CostTracker()
    build_set([_niche("cosy mystery a")], tmp_path / "v.xlsx",
              fetch_shelf=faux_shelf, classify=faux_classify, cost=cost)
    assert cost.breakdown()["llm_tokens_in"] == 1000
    assert cost.breakdown()["llm_usd"] > 0
