"""Outillage partagé des tests de la CLI de calibration low-content (groupe 4).

Données RÉELLES d'abord, et chaque écart est DÉCLARÉ :
- `v2_asin_payloads.json` : 8 captures réelles amazon.fr du 2026-07-20, parsées par le
  parseur COURANT. `sans_pages=True` retire la pagination APRÈS parsing pour reproduire une
  fiche lue par l'ancien parseur (run 4) : c'est une ALTÉRATION, pas une capture.
- Les SERP construites ici sont INVENTÉES dans leur forme (rang, ordre, rotation des 8 ASIN
  réels) : aucune SERP brute n'est versionnée. Titres et prix viennent des captures.
- `99-logs/validation-lc.xlsx` : le classeur réel de Baptiste (versionné), LU seulement.
- `cache_run4_reel` : l'état RÉEL du cache après le run 4 pour les 31 requêtes du classeur
  (31 SERP, 185 fiches, 84 classifications), extrait en lecture seule le 2026-09-14 dans
  `fixtures/cache_calibration_run4_reel.json` — seule altération déclarée : organiques
  tronqués aux 12 premiers (le jeu en lit 6). Versionné : `99-logs/df-cache.db` est ignoré
  par git, et les tests qui le copiaient étaient IGNORÉS sur tout clone neuf, donc « 0
  ignoré » y était impossible. Les expirations sont posées relativement à maintenant, pour
  que le test voie le cache tel qu'il était le 2026-09-14 à 12:00 — il ne se périme pas.
"""
from __future__ import annotations

import json
import sqlite3
import time
from pathlib import Path

from cache import Cache
from fiction_books import parse_enriched_book
from models import SearchItem, SearchResult

RACINE = Path(__file__).resolve().parent.parent
XLSX_REEL = RACINE / "99-logs" / "validation-lc.xlsx"
PAYLOADS = json.loads((Path(__file__).parent / "fixtures" / "fiction" /
                       "v2_asin_payloads.json").read_text(encoding="utf-8"))
ASINS_V2 = list(PAYLOADS)
LOC, LANG = 2250, "fr_FR"


def etiquetees_reelles():
    from lowcontent_validation import charger_etiquettes
    return charger_etiquettes(XLSX_REEL)


# Le classeur a grandi le 2026-09-15 : 20 requêtes NEUVES (colonne E « lot » = « neuf
# 2026-09-15 ») à côté des 31 du run 4 (« run 4 »). Les tests qui photographient l'état du
# cache APRÈS le run 4 (fixture `cache_calibration_run4_reel.json`, 31 SERP, 185 fiches) ne
# portent que sur ces 31-là : le classeur réel continue de vivre, la photo du run 4 non.
LOT_RUN4 = "run 4"


def _requetes_du_lot(lot: str) -> set[str]:
    import openpyxl
    ws = openpyxl.load_workbook(XLSX_REEL, read_only=True).active
    return {str(r[0]).strip() for r in ws.iter_rows(values_only=True)
            if r and len(r) > 4 and r[0] and str(r[4] or "").strip() == lot}


def etiquetees_run4():
    """Les 31 requêtes étiquetées du run 4, lues dans le classeur RÉEL (colonne « lot »)."""
    run4 = _requetes_du_lot(LOT_RUN4)
    return [e for e in etiquetees_reelles() if e.requete in run4]


def xlsx_run4(dossier: Path) -> Path:
    """Copie du classeur réel où seules restent les lignes du lot « run 4 » (les autres
    lignes de données sont VIDÉES, ce que `charger_etiquettes` ignore) — pour les tests qui
    passent un classeur à la CLI. Le classeur réel n'est jamais écrit."""
    import openpyxl
    from lowcontent_validation import _ENTETES
    wb = openpyxl.load_workbook(XLSX_REEL)
    ws = wb.active
    apres_entete = False
    for ligne in ws.iter_rows():
        valeurs = [str(c.value or "").strip().lower() for c in ligne[:4]]
        if valeurs == list(_ENTETES):
            apres_entete = True
            continue
        if apres_entete and str(ligne[4].value if len(ligne) > 4 else "").strip() != LOT_RUN4:
            for c in ligne[:5]:
                c.value = None
    dest = Path(dossier) / "validation-lc-run4.xlsx"
    wb.save(dest)
    return dest


def livre_v2(asin: str, sans_pages: bool = False):
    b = parse_enriched_book(PAYLOADS[asin])
    return b.model_copy(update={"pages": None}) if sans_pages else b


def serp_v2(requete: str, asins: list[str]) -> SearchResult:
    organic = []
    for i, a in enumerate(asins, 1):
        b = livre_v2(a)
        organic.append(SearchItem(rank=i, asin=a, title=b.title, price=b.price))
    return SearchResult(keyword=requete, organic=organic, total_items=len(organic))


def asins_de(i: int, n: int = 6) -> list[str]:
    """Rotation des 8 ASIN réels : deux requêtes voisines partagent des fiches, comme au
    run 4 (186 ASIN pour 185 fiches)."""
    return [ASINS_V2[(i + k) % len(ASINS_V2)] for k in range(n)]


def cache_seme(chemin: Path, requetes: list[str], sans_pages: bool = True,
               serp: bool = True, livres: bool = True) -> Cache:
    c = Cache(chemin)
    for i, q in enumerate(requetes):
        asins = asins_de(i)
        if serp:
            c.set_search(q, LOC, LANG, serp_v2(q, asins), 15 * 24 * 3600)
        if livres:
            for a in asins:
                c.set_book(a, LOC, livre_v2(a, sans_pages=sans_pages), 15 * 24 * 3600)
    return c


FIXTURE_RUN4 = Path(__file__).parent / "fixtures" / "cache_calibration_run4_reel.json"


def cache_run4_reel(tmp_path: Path, nom: str = "cache-run4.db") -> Path:
    """Reconstruit le cache du run 4. Les clés sont recalculées par `Cache` (et non recopiées) :
    une empreinte de schéma qui change plus tard ne fait pas lire la fixture comme vide."""
    donnees = json.loads(FIXTURE_RUN4.read_text(encoding="utf-8"))
    dst = tmp_path / nom
    Cache(dst)
    maintenant = time.time()
    lignes = [(Cache._search_key(e["requete"], LOC, LANG), e["value"], e["expires_moins_ref"])
              for e in donnees["search"]]
    lignes += [(Cache._book_key(e["asin"], LOC), e["value"], e["expires_moins_ref"])
               for e in donnees["book"]]
    lignes += [(e["key"], e["value"], e["expires_moins_ref"]) for e in donnees["clf"]]
    cx = sqlite3.connect(dst)
    try:
        cx.executemany("INSERT OR REPLACE INTO kv (key, value, expires) VALUES (?, ?, ?)",
                       [(k, json.dumps(v, ensure_ascii=False), maintenant + d)
                        for k, v, d in lignes])
        cx.commit()
    finally:
        cx.close()
    return dst


def cles(chemin: Path) -> dict[str, tuple[str, float]]:
    cx = sqlite3.connect(chemin)
    try:
        return {k: (v, e) for k, v, e in cx.execute("SELECT key, value, expires FROM kv")}
    finally:
        cx.close()
